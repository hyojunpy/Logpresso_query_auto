from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from scripts.store_field_metadata import enrich_payload


TYPE_MAP = {
    "DATE": "datetime", "LONG": "integer", "INT": "integer", "INTEGER": "integer",
    "DOUBLE": "float", "FLOAT": "float", "IP": "ip", "BOOL": "boolean",
    "BOOLEAN": "boolean", "PORT": "integer", "STRING": "string", "COUNTRY": "string",
    "ARRAY": "array", "미표기": "unknown",
}

PRODUCT_ALIASES = {
    "BLUEMAX NGF": ["블루맥스 NGF", "NGF", "차세대 방화벽"],
    "SECUI MF2": ["시큐아이 MF2", "MF2"], "SECUI MFD": ["시큐아이 MFD", "MFD"],
    "BLUEMAX IPS": ["블루맥스 IPS"], "BLUEMAX ADS": ["블루맥스 ADS"],
    "BLUEMAX WIPS": ["블루맥스 WIPS"], "SECUI MFI": ["시큐아이 MFI", "MFI"],
    "FortiGate": ["포티게이트", "Fortinet FortiGate"],
    "AhnLab TrusGuard": ["안랩 트러스가드", "TrusGuard", "트러스가드"],
    "AhnLab AIPS": ["안랩 AIPS"], "AhnLab DPX": ["안랩 DPX"],
    "AhnLab MDS": ["안랩 MDS"], "Cisco Meraki": ["시스코 머라키", "Meraki", "머라키"],
    "Palo Alto Networks NGFW": ["Palo Alto NGFW", "팔로알토 NGFW"],
    "Genian NAC": ["지니안 NAC"], "Genian EDR": ["지니안 EDR"],
    "QueryPie DAC": ["쿼리파이 DAC"], "SNIPER ONE-i": ["스나이퍼 ONE-i"],
    "SNIPER TMS Plus": ["스나이퍼 TMS Plus"], "SNIPER APTX": ["스나이퍼 APTX"],
}

FIELD_ALIASES = {
    "_time": ["시각", "시간", "발생 시간", "발생시간", "이벤트 시간", "로그 시간", "일시"],
    "hostname": ["호스트", "호스트명", "장비명", "시스템명"],
    "device_name": ["장비명", "디바이스명", "기기명"],
    "src_ip": ["출발지 IP", "출발지IP", "출발지 주소", "소스 IP", "송신지 IP", "공격자 IP", "클라이언트 IP"],
    "dst_ip": ["목적지 IP", "목적지IP", "도착지 IP", "대상 IP", "수신지 IP", "서버 IP"],
    "src_port": ["출발지 포트", "출발지포트", "소스 포트", "송신 포트", "클라이언트 포트"],
    "dst_port": ["목적지 포트", "목적지포트", "도착지 포트", "대상 포트", "서비스 포트", "서버 포트"],
    "action": ["행위", "동작", "조치", "처리 결과", "처리결과", "차단 여부", "정책 행위"],
    "protocol": ["프로토콜", "통신 규약"], "user": ["사용자", "사용자명", "계정", "아이디"],
    "user_name": ["사용자 이름", "사용자명", "계정명"],
    "bytes": ["바이트", "전송량", "트래픽량", "데이터량"],
    "tx_bytes": ["송신 바이트", "송신바이트수", "보낸 바이트", "송신량"],
    "rx_bytes": ["수신 바이트", "수신바이트수", "받은 바이트", "수신량"],
    "tx_pkts": ["송신 패킷", "송신패킷수", "보낸 패킷"],
    "rx_pkts": ["수신 패킷", "수신패킷수", "받은 패킷"],
    "total_bytes": ["전체 바이트", "총 바이트", "누적 바이트", "트래픽 바이트"],
    "total_pkts": ["전체 패킷", "총 패킷", "누적 패킷"],
    "cpu_usage": ["CPU 사용률", "CPU 점유율", "프로세서 사용률"],
    "mem_usage": ["메모리 사용률", "메모리 점유율", "RAM 사용률"],
    "disk_usage": ["디스크 사용률", "저장공간 사용률"],
    "domain": ["도메인", "접속 도메인"], "url": ["URL", "웹 주소", "접속 주소"],
    "dns_query": ["DNS 질의", "질의 도메인", "조회 도메인", "도메인"],
    "duration": ["실행 시간", "수행 시간", "소요 시간", "처리 시간"],
    "severity": ["심각도", "위험도", "위협 등급"], "risk": ["위험도", "위험 등급"],
    "risk_score": ["위험도 점수", "위험 점수", "위협 점수"],
    "status": ["상태", "처리 상태"], "result": ["결과", "처리 결과"],
    "rule_id": ["룰 ID", "규칙 ID", "정책 ID"], "policy": ["정책", "보안 정책"],
    "signature": ["시그니처", "탐지명", "공격명"], "msg": ["메시지", "로그 메시지", "내용"],
    "line": ["원문", "원본 로그", "시스로그 메시지"],
}

FIELD_TOKEN_ALIASES = {
    "account": ["계정", "계정명", "사용자 계정"],
    "client": ["클라이언트", "접속자"],
    "server": ["서버", "대상 서버"],
    "remote": ["원격", "원격지"],
    "local": ["로컬", "내부"],
    "source": ["출발지", "소스", "송신지"],
    "destination": ["목적지", "도착지", "대상"],
    "src": ["출발지", "소스", "송신지"],
    "dst": ["목적지", "도착지", "대상"],
    "ip": ["IP", "주소"],
    "port": ["포트"],
    "username": ["사용자명", "계정명", "아이디"],
    "userid": ["사용자 ID", "계정 ID"],
    "query": ["질의", "쿼리"],
    "sql": ["SQL", "쿼리"],
    "elapsed": ["경과 시간", "소요 시간", "처리 시간"],
    "latency": ["지연 시간", "응답 시간"],
    "event": ["이벤트"],
    "category": ["분류", "카테고리"],
    "method": ["메서드", "요청 방식"],
    "path": ["경로", "요청 경로"],
}

LOG_TYPE_ALIASES = {
    "urlblock": ["웹 차단", "URL 차단", "유해 사이트 차단", "웹 필터링"],
    "sslvpn_user_auth": ["SSL VPN 인증", "VPN 사용자 인증", "원격 접속 인증", "VPN 로그인"],
    "sslvpn_monitoring": ["SSL VPN 상태", "VPN 터널 상태", "원격 접속 상태"],
    "ipsecvpn_event": ["IPSEC VPN 이벤트", "IPSEC 이벤트", "사이트간 VPN 이벤트"],
    "ips_ddos_detect": ["침입 탐지", "공격 탐지", "DDoS 탐지", "디도스 탐지"],
    "ips_ddos_incident": ["공격 통계", "침해 통계", "DDoS 통계", "디도스 통계"],
    "dns_security": ["DNS 보안", "악성 도메인", "DNS 위협"],
    "resource_cnt": ["시스템 성능", "장비 성능", "리소스 사용량", "자원 사용률"],
    "if_traffic_cnt": ["인터페이스 트래픽", "인터페이스 사용량", "회선 트래픽"],
    "ha_event": ["HA 이벤트", "이중화 이벤트", "절체 이벤트"],
    "ha_status_cnt": ["HA 상태", "이중화 상태", "절체 상태"],
    "fw4_blacklist": ["블랙리스트", "차단 목록", "금지 IP 목록"],
}

# These Store parsers expose the schema code as log_type even when some schema
# pages omit that shared discriminator from their per-schema field tables.
LOG_TYPE_PRODUCTS = {
    "BLUEMAX NGF", "SECUI MF2", "SECUI MFD", "BLUEMAX IPS",
    "BLUEMAX ADS", "BLUEMAX WIPS", "SECUI MFI",
}


def _cell(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _codes(value: str) -> list[str]:
    return [item.strip() for item in re.split(r"[,/]", value) if item.strip()]


def _product_aliases(manufacturer: str, product: str) -> list[str]:
    aliases = {product, *PRODUCT_ALIASES.get(product, [])}
    aliases.add(re.sub(r"\s*\([^)]*\)\s*$", "", product).strip())
    if "/" in product:
        aliases.update(part.strip() for part in product.split("/") if len(part.strip()) >= 3)
    if manufacturer and product.lower().startswith(manufacturer.lower() + " "):
        aliases.add(product[len(manufacturer):].strip())
    return sorted((item for item in aliases if len(item) >= 2), key=lambda item: (-len(item), item))


def _schema_aliases(product_aliases: list[str], schema_name: str, log_types: list[str]) -> tuple[list[str], list[str]]:
    short_aliases: set[str] = set()
    for prefix in product_aliases:
        stripped = re.sub(rf"^{re.escape(prefix)}\s*", "", schema_name, flags=re.IGNORECASE).strip()
        if stripped != schema_name and len(stripped) >= 2:
            short_aliases.add(stripped)
    product_words = product_aliases[0].split()
    schema_words = schema_name.split()
    common = 0
    for product_word, schema_word in zip(product_words, schema_words):
        if product_word.lower() != schema_word.lower():
            break
        common += 1
    if common:
        stripped = " ".join(schema_words[common:]).strip()
        if len(stripped) >= 2:
            short_aliases.add(stripped)
    for log_type in log_types:
        short_aliases.update(LOG_TYPE_ALIASES.get(log_type, []))
    return [schema_name], sorted(short_aliases, key=lambda item: (-len(item), item))


def _derived_field_aliases(name: str) -> set[str]:
    """Build conservative Korean aliases from common field-name tokens."""
    aliases: set[str] = set()
    compact = name.lower().replace("_", "")
    aliases.update(FIELD_TOKEN_ALIASES.get(compact, []))
    tokens = [token for token in re.split(r"[_-]+", name.lower()) if token]
    if len(tokens) == 2:
        left = FIELD_TOKEN_ALIASES.get(tokens[0], [])
        right = FIELD_TOKEN_ALIASES.get(tokens[1], [])
        for left_alias in left[:3]:
            for right_alias in right[:2]:
                aliases.add(f"{left_alias} {right_alias}")
                aliases.add(f"{left_alias}{right_alias}")
    return aliases


def _field(name: str, source_type: str, display_name: str, description: str) -> dict[str, object]:
    aliases = set(FIELD_ALIASES.get(name, []))
    aliases.update(_derived_field_aliases(name))
    if display_name:
        aliases.update({display_name, re.sub(r"\s+", "", display_name)})
    return {
        "name": name, "type": TYPE_MAP.get(source_type.upper(), "unknown"),
        "source_type": source_type, "display_name": display_name, "description": description,
        "aliases": sorted((item for item in aliases if len(item) >= 2), key=lambda item: (-len(item), item)),
    }


def build(source: Path) -> dict[str, object]:
    # Load from memory so openpyxl never holds a Windows lock on the upload
    # temporary file while the caller removes it.
    workbook = load_workbook(BytesIO(source.read_bytes()), read_only=True, data_only=True)
    version = "unknown"
    if "요약" in workbook.sheetnames:
        for row in workbook["요약"].iter_rows(values_only=True):
            if _cell(row[0] if row else "") == "기준일":
                version = _cell(row[1] if len(row) > 1 else "") or version
                break
    metadata: dict[str, dict[str, object]] = {}
    for row in workbook["제품_목록"].iter_rows(min_row=2, values_only=True):
        manufacturer, product = _cell(row[0]), _cell(row[1])
        if product:
            metadata[product] = {
                "manufacturer": manufacturer, "verification": _cell(row[2]), "protocol": _cell(row[3]),
                "port": _cell(row[4]), "format": _cell(row[5]), "scope": _cell(row[6]), "source_url": _cell(row[8]),
            }

    aliases_by_product = {
        product: _product_aliases(str(meta.get("manufacturer") or ""), product)
        for product, meta in metadata.items()
    }
    schemas: dict[tuple[str, str], dict[str, object]] = {}
    for row in workbook["스키마_목록"].iter_rows(min_row=2, values_only=True):
        manufacturer, product, schema_name = _cell(row[0]), _cell(row[1]), _cell(row[2])
        if _cell(row[8]) != "실제" or not product or not schema_name:
            continue
        metadata.setdefault(product, {"manufacturer": manufacturer})
        product_aliases = aliases_by_product.setdefault(product, _product_aliases(manufacturer, product))
        log_types = _codes(_cell(row[3]))
        aliases, short_aliases = _schema_aliases(product_aliases, schema_name, log_types)
        schemas[(product, schema_name)] = {
            "name": schema_name, "aliases": aliases, "short_aliases": short_aliases,
            "log_types": log_types, "status": _cell(row[5]), "source_url": _cell(row[6]),
            "note": _cell(row[7]), "fields": [], "discriminator_field": None,
        }

    fields_by_schema: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in workbook["필드_상세"].iter_rows(min_row=2, values_only=True):
        product, schema_name, field_name = _cell(row[1]), _cell(row[2]), _cell(row[5])
        if product and schema_name and field_name:
            fields_by_schema[(product, schema_name)].append(
                _field(field_name, _cell(row[6]), _cell(row[7]), _cell(row[8]))
            )
    for key, schema in schemas.items():
        fields = fields_by_schema.get(key, [])
        schema["fields"] = fields
        names = {str(field["name"]) for field in fields}
        schema["discriminator_field"] = next((name for name in ("log_type", "module_flag") if name in names), None)
        if schema["discriminator_field"] is None and key[0] in LOG_TYPE_PRODUCTS and schema["log_types"]:
            schema["discriminator_field"] = "log_type"

    raw_formats: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in workbook["Raw_양식_공개"].iter_rows(min_row=2, values_only=True):
        product = _cell(row[0])
        if product:
            raw_formats[product].append({
                "log_type": _cell(row[1]), "transport": _cell(row[2]),
                "template": _cell(row[3]), "source_url": _cell(row[4]),
            })

    common_fields = [
        _field(_cell(row[0]), _cell(row[1]), "", _cell(row[2]))
        for row in workbook["공통_Syslog_필드"].iter_rows(min_row=2, values_only=True) if _cell(row[0])
    ]
    schemas_by_product: dict[str, list[dict[str, object]]] = defaultdict(list)
    for (product, _), schema in schemas.items():
        schemas_by_product[product].append(schema)
    products = []
    for product, product_schemas in sorted(schemas_by_product.items()):
        meta = metadata.get(product, {})
        products.append({
            "name": product, "manufacturer": meta.get("manufacturer", ""),
            "aliases": aliases_by_product.get(product, [product]),
            "verification": meta.get("verification", ""), "protocol": meta.get("protocol", ""),
            "port": meta.get("port", ""), "format": meta.get("format", ""), "scope": meta.get("scope", ""),
            "source_url": meta.get("source_url", ""), "raw_formats": raw_formats.get(product, []),
            "schemas": sorted(product_schemas, key=lambda item: str(item["name"])),
        })
    stats = {
        "products": len(products), "schemas": len(schemas),
        "schemas_with_fields": sum(bool(schema["fields"]) for schema in schemas.values()),
        "fields": sum(len(schema["fields"]) for schema in schemas.values()),
        "raw_formats": sum(len(items) for items in raw_formats.values()),
    }
    workbook.close()
    return enrich_payload({
        "version": version, "source": source.name, "stats": stats,
        "common_fields": common_fields, "products": products,
    })


def main() -> None:
    parser = argparse.ArgumentParser(description="Build Logpresso Store Syslog schema knowledge JSON")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    payload = build(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    stats = payload["stats"]
    print(f"Wrote {stats['products']} products, {stats['schemas']} schemas, {stats['fields']} fields and {stats['raw_formats']} raw formats to {args.output}")


if __name__ == "__main__":
    main()
