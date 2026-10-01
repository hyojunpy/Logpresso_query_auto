from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


EXAMPLES = [
    (
        "라이선스 만료일",
        "라이선스 만료일과 만료 임박 상태를 조회합니다.",
        '''confdb docs logpresso-license licenses
| parsekv field=content kvdelim="=" overlay=t pairdelim="\\n"
| fields hardware_key, installed, not_after, site, support_expiry, volume_limit
| rename site as 사이트, installed as 갱신날짜, support_expiry as 유지보수기한, hardware_key as 하드웨어키, not_after as 라이선스만료일
| eval 라이선스종류=case(contains(사이트, "dev"), "개발", contains(사이트, "tmp"), "임시", "정식")
| eval 라이선스만료일=date(라이선스만료일, "yyyy-MM-dd hh:mm:ss")
| eval date_diff=datediff(now(), 라이선스만료일, "day")
| eval text=case(date_diff == 1, "HIGH(라이선스 1일 뒤 만료)", date_diff == 7, "MIDDLE(라이선스 7일 뒤 만료)", date_diff == 30, "LOW(라이선스 30일 뒤 만료)", "정상")
| eval 라이선스만료일=string(라이선스만료일, "yyyy-MM-dd")
| limit 10000''',
    ),
    (
        "일 로그 수집 현황",
        "지정한 하루 범위의 전체 수집량과 로그 건수를 조회합니다.",
        '''set from=$("from", ago("1d"))
| set to=$("_to", now())
| table from=$("from") to=$("to") *:sys_logger_stats
| search _node != "forwarder_1"
| stats sum(volume) as volume, sum(count) as count
| eval volume=round(volume / 1024 / 1024)
| eval count=format("%,d", count)
| eval count=concat(count, "건")
| limit 10000''',
    ),
    (
        "라이선스 종류",
        "사이트명 규칙을 기준으로 개발 임시 정식 라이선스를 분류합니다.",
        '''confdb docs logpresso-license licenses
| parsekv field=content kvdelim="=" overlay=t pairdelim="\\n"
| fields hardware_key, installed, not_after, site, support_expiry, volume_limit
| rename site as 사이트, installed as 갱신날짜, support_expiry as 유지보수기한, hardware_key as 하드웨어키, not_after as 라이선스만료일
| eval 라이선스종류=case(contains(사이트, "dev"), "개발", contains(사이트, "tmp"), "임시", "정식")
| limit 10000''',
    ),
    (
        "수집 속도",
        "포워더의 최신 초당 수집 속도를 계산합니다.",
        '''table limit=1 *:sys_logger_trends
| eval rate=round(delta * 1000 / interval)
| search _node == "forwarder"
| fields rate
| limit 10000''',
    ),
    (
        "라이선스 사용 추이",
        "노드별 라이선스 사용량을 일 단위 시계열로 집계합니다.",
        '''system license-usages
| timechart span=1d sum(volume) as volume by node
| limit 10000''',
    ),
    (
        "최근 한 달 라이선스 초과 사용일",
        "일별 사용량을 계약 용량과 비교해 초과 상태를 표시합니다.",
        '''system license-usages
| timechart span=1d sum(volume) as volume
| eval contract_volume=10737418240
| eval 초과상태=if(volume > contract_volume, "라이선스 초과 사용", "정상")
| eval contract_volume=concat(round(contract_volume / 1024 / 1024 / 1024, 2), "GB")
| eval volume=concat(round(volume / 1024 / 1024 / 1024, 2), "GB")
| rename _time as 날짜, volume as 사용량, contract_volume as 계약용량
| order 날짜, 사용량, 초과상태, 계약용량
| limit 10000''',
    ),
    (
        "최근 7일 라이선스 사용량 상위 수집기",
        "최근 7일간 수집기별 사용량을 집계하고 수집기 이름을 결합합니다.",
        '''table duration=7d *:sys_logger_stats
| stats sum(volume) as volume by logger_name
| rex field=logger_name "sonar_logger_(?<id>\\d{5})"
| eval id=int(id)
| join type=left id [dbquery sonar select id, name from sonar_loggers]
| sort -volume
| limit 7
| eval 순번=seq()
| rename name as 수집기명, volume as 수집량
| fields 순번, 수집기명, 수집량
| search isnotnull(수집기명)
| eval len_volume=len(수집량)
| eval 수집량=case(len_volume < 10, concat(round(수집량 / 1024 / 1024, 2), "MB"), concat(round(수집량 / 1024 / 1024 / 1024, 2), "GB"))
| fields -len_volume
| limit 10000''',
    ),
    (
        "전일 수집기별 수집량",
        "전일 수집량에 수집기와 자산 정보를 결합하고 읽기 쉬운 단위로 표시합니다.",
        '''set to=datetrunc(now(), "1d")
| set from=datetrunc(dateadd($("to"), "day", -1), "1d")
| table from=$("from") to=$("to") *:sys_logger_stats
| eval _time=datetrunc(_time, "1d")
| stats sum(volume) as volume by logger_name, _time
| search logger_name != "*sonar_system_alert*"
| eval id=int(valueof(groups(logger_name, "sonar_logger_(\\d+)"), 0))
| join type=left id [sonar loggers | fields asset_ip, hostname, id, name, table_name]
| search isnotnull(table_name)
| eval len_volume=len(volume)
| sort -volume
| eval volume=case(len_volume <= 3, concat(volume, "B"), len_volume >= 4 and len_volume < 7, concat(round(volume / 1024, 2), "KB"), len_volume >= 7 and len_volume < 10, concat(round(volume / 1024 / 1024, 2), "MB"), len_volume >= 10, concat(round(volume / 1024 / 1024 / 1024, 2), "GB"))
| fields -logger_name, id, len_volume
| order _time, name, table_name, volume, asset_ip, hostname
| eval asset_ip=if(isnull(asset_ip), "자산IP 미등록", asset_ip)
| eval hostname=if(isnull(hostname), "자산IP 미등록", hostname)
| limit 10000''',
    ),
]


def shade(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F5F7")
    properties.append(shading)


def append_examples(path: Path) -> None:
    document = Document(path)
    if any(paragraph.text == "운영 대시보드 쿼리 예제" for paragraph in document.paragraphs):
        return
    document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("운영 대시보드 쿼리 예제", level=1)
    document.add_paragraph(
        "로그프레소 운영 현황 대시보드에서 반복적으로 사용하는 라이선스와 로그 수집 지표의 자연어 의도 및 쿼리 패턴입니다. "
        "제목과 설명은 검색 의도로 사용하고 쿼리는 명령 순서와 계산식을 유지해 생성합니다."
    )
    for title, description, query in EXAMPLES:
        document.add_heading(title, level=2)
        document.add_paragraph(description)
        paragraph = document.add_paragraph()
        shade(paragraph)
        for index, line in enumerate(query.splitlines()):
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
            run.font.size = Pt(8)
            if index < len(query.splitlines()) - 1:
                run.add_break()
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
    document.save(path)


if __name__ == "__main__":
    append_examples(Path("docs") / "로그프레소 쿼리.docx")
