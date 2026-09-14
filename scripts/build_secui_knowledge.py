from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


TYPE_MAP = {
    "DATE": "datetime",
    "LONG": "integer",
    "INT": "integer",
    "DOUBLE": "float",
    "IP": "ip",
    "BOOL": "boolean",
    "PORT": "integer",
    "STRING": "string",
    "COUNTRY": "string",
    "ARRAY": "array",
    "미표기": "unknown",
}

PRODUCT_ALIASES = {
    "BLUEMAX NGF": ["BLUEMAX NGF", "블루맥스 NGF", "NGF", "차세대 방화벽"],
    "SECUI MF2": ["SECUI MF2", "시큐아이 MF2", "MF2"],
    "SECUI MFD": ["SECUI MFD", "시큐아이 MFD", "MFD"],
    "BLUEMAX IPS": ["BLUEMAX IPS", "블루맥스 IPS", "IPS", "침입방지시스템"],
    "BLUEMAX ADS": ["BLUEMAX ADS", "블루맥스 ADS", "ADS", "이상징후 탐지"],
    "BLUEMAX WIPS": ["BLUEMAX WIPS", "블루맥스 WIPS", "WIPS", "무선 침입방지"],
    "SECUI MFI": ["SECUI MFI", "시큐아이 MFI", "MFI"],
}

FIELD_ALIASES = {
    "_time": ["시각", "시간", "발생 시간", "발생시간", "이벤트 시간", "로그 시간", "일시"],
    "hostname": ["호스트", "호스트명", "장비명", "시스템명"],
    "src_ip": ["출발지 IP", "출발지IP", "출발지 주소", "소스 IP", "송신지 IP", "공격자 IP"],
    "dst_ip": ["목적지 IP", "목적지IP", "도착지 IP", "대상 IP", "수신지 IP", "타깃 IP"],
    "src_port": ["출발지 포트", "소스 포트", "송신 포트"],
    "dst_port": ["목적지 포트", "도착지 포트", "대상 포트", "서비스 포트"],
    "action": ["행위", "동작", "처리 결과", "처리결과", "차단 여부", "정책 행위"],
    "protocol": ["프로토콜", "통신 규약"],
    "user": ["사용자", "사용자명", "계정"],
    "bytes": ["바이트", "전송량", "트래픽량", "데이터량"],
    "tx_bytes": ["송신 바이트", "송신바이트수", "보낸 바이트", "송신량"],
    "rx_bytes": ["수신 바이트", "수신바이트수", "받은 바이트", "수신량"],
    "tx_pkts": ["송신 패킷", "송신패킷수", "보낸 패킷"],
    "rx_pkts": ["수신 패킷", "수신패킷수", "받은 패킷"],
    "total_bytes": ["전체 바이트", "총 바이트", "누적 바이트", "트래픽 바이트"],
    "total_pkts": ["전체 패킷", "총 패킷", "누적 패킷"],
    "cpu_usage": ["CPU 사용률", "CPU 점유율", "프로세서 사용률"],
    "mem_usage": ["메모리 사용률", "메모리 점유율", "RAM 사용률"],
    "disk_usage": ["디스크 사용률", "저장공간 사용률"],
    "domain": ["도메인", "접속 도메인"],
    "url": ["URL", "웹 주소", "접속 주소"],
    "severity": ["심각도", "위험도", "위협 등급"],
    "status": ["상태", "처리 상태"],
    "rule_id": ["룰 ID", "규칙 ID", "정책 ID"],
    "policy": ["정책", "보안 정책"],
}

LOG_TYPE_ALIASES = {
    "urlblock": ["웹 차단", "URL 차단", "유해 사이트 차단", "웹 필터링"],
    "sslvpn_user_auth": ["SSL VPN 인증", "VPN 사용자 인증", "원격 접속 인증", "VPN 로그인"],
    "sslvpn_monitoring": ["SSL VPN 상태", "VPN 터널 상태", "원격 접속 상태"],
    "ipsecvpn_event": ["IPSEC VPN 이벤트", "IPSEC 이벤트", "사이트간 VPN 이벤트"],
    "ips_ddos_detect": ["침입 탐지", "공격 탐지", "DDoS 탐지", "디도스 탐지"],
    "ips_ddos_incident": ["공격 통계", "침해 통계", "DDoS 통계", "디도스 통계"],
    "dns_security": ["DNS 보안", "악성 도메인", "DNS 위협"],
    "resource_cnt": ["시스템 성능", "장비 성능", "리소스 사용량", "자원 사용률"],
    "if_traffic_cnt": ["인터페이스 트래픽", "인터페이스 사용량", "회선 트래픽"],
    "ha_event": ["HA 이벤트", "이중화 이벤트", "절체 이벤트"],
    "ha_status_cnt": ["HA 상태", "이중화 상태", "절체 상태"],
    "fw4_blacklist": ["블랙리스트", "차단 목록", "금지 IP 목록"],
}


def _cell(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _schema_aliases(product: str, schema_name: str, log_types: list[str]) -> list[str]:
    aliases = {schema_name}
    stripped = schema_name
    for prefix in (product, "블루맥스 NGF", "블루맥스 IPS", "BLUEMAX ADS", "MFD", "MF2"):
        stripped = re.sub(rf"^{re.escape(prefix)}\s*", "", stripped, flags=re.IGNORECASE)
    if len(stripped) >= 2:
        aliases.add(stripped)
    for log_type in log_types:
        aliases.update(LOG_TYPE_ALIASES.get(log_type, []))
    return sorted(aliases, key=lambda value: (-len(value), value))


def build(source: Path) -> dict[str, object]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    schema_rows: dict[tuple[str, str], dict[str, object]] = {}
    for row in workbook["스키마_목록"].iter_rows(min_row=2, values_only=True):
        product, schema_name, log_types = _cell(row[1]), _cell(row[2]), _cell(row[3])
        if not product or not schema_name:
            continue
        parsed_log_types = [item.strip() for item in log_types.split(",") if item.strip()]
        schema_rows[(product, schema_name)] = {
            "name": schema_name,
            "aliases": _schema_aliases(product, schema_name, parsed_log_types),
            "log_types": parsed_log_types,
            "status": _cell(row[5]),
            "format_note": _cell(row[6]),
            "source_url": _cell(row[8]),
            "fields": [],
        }

    field_rows: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in workbook["필드_상세"].iter_rows(min_row=2, values_only=True):
        product, schema_name, field_name = _cell(row[0]), _cell(row[1]), _cell(row[4])
        if not product or not schema_name or not field_name:
            continue
        display_name = _cell(row[6])
        aliases = set(FIELD_ALIASES.get(field_name, []))
        if display_name:
            aliases.update({display_name, re.sub(r"\s+", "", display_name)})
        field_rows[(product, schema_name)].append(
            {
                "name": field_name,
                "type": TYPE_MAP.get(_cell(row[5]).upper(), "unknown"),
                "source_type": _cell(row[5]),
                "display_name": display_name,
                "description": _cell(row[7]),
                "aliases": sorted((alias for alias in aliases if len(alias) >= 2), key=lambda value: (-len(value), value)),
            }
        )

    products: dict[str, list[dict[str, object]]] = defaultdict(list)
    for key, schema in schema_rows.items():
        schema["fields"] = field_rows.get(key, [])
        products[key[0]].append(schema)
    return {
        "version": "2026-09-14",
        "source": source.name,
        "products": [
            {
                "name": product,
                "aliases": PRODUCT_ALIASES.get(product, [product]),
                "schemas": sorted(schemas, key=lambda item: str(item["name"])),
            }
            for product, schemas in sorted(products.items())
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Build SECUI schema knowledge JSON from the reviewed workbook")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    payload = build(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    schema_count = sum(len(product["schemas"]) for product in payload["products"])
    field_count = sum(len(schema["fields"]) for product in payload["products"] for schema in product["schemas"])
    print(f"Wrote {schema_count} schemas and {field_count} fields to {args.output}")


if __name__ == "__main__":
    main()
